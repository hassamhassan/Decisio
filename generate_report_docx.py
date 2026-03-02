"""
Generate Decisio Flow Report as DOCX
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# ── Styles ──────────────────────────────────────────────────────────
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

style_h1 = doc.styles['Heading 1']
style_h1.font.size = Pt(22)
style_h1.font.color.rgb = RGBColor(0x1a, 0x56, 0xdb)

style_h2 = doc.styles['Heading 2']
style_h2.font.size = Pt(16)
style_h2.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf)

style_h3 = doc.styles['Heading 3']
style_h3.font.size = Pt(13)
style_h3.font.color.rgb = RGBColor(0x37, 0x30, 0xa3)


def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for row_data in rows:
        row = t.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
            for p in row.cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def p(text, bold=False, italic=False, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return para


def bullet(text, level=0):
    para = doc.add_paragraph(text, style='List Bullet')
    para.paragraph_format.left_indent = Inches(0.25 * (level + 1))


# ── Title Page ──────────────────────────────────────────────────────
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_heading('Decisio', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(0x1a, 0x56, 0xdb)

subtitle = doc.add_heading('Application Flow & Architecture Report', level=0)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in subtitle.runs:
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

doc.add_paragraph()
p('Operational Decision Support System', italic=True, size=14).alignment = WD_ALIGN_PARAGRAPH.CENTER
p('February 2026', italic=True, size=12).alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ── 1. System Overview ──────────────────────────────────────────────
doc.add_heading('1. System Overview', level=1)
p('Decisio is an operational decision-support system built for industrial environments. '
  'It guides operators through structured incident diagnosis using an AI-driven 10-step '
  'diagnostic framework, then generates actionable Decision Briefs with risk-rated options.')
p('')
p('The system captures why decisions are made, not how to execute repairs.', bold=True, italic=True)

# ── 2. Tech Stack ───────────────────────────────────────────────────
doc.add_heading('2. Technology Stack', level=1)
add_table(
    ['Layer', 'Technology'],
    [
        ['Frontend', 'React 19 + Vite, React Router'],
        ['Backend', 'FastAPI (Python 3.11), async + sync'],
        ['AI Pipeline', 'LangGraph (StateGraph), LangChain'],
        ['LLM', 'Azure OpenAI (GPT-4o)'],
        ['Database', 'PostgreSQL (async via SQLAlchemy)'],
        ['Vector Store', 'Qdrant (Decision Memory)'],
        ['Auth', 'JWT + bcrypt'],
    ]
)

# ── 3. Architecture ─────────────────────────────────────────────────
doc.add_heading('3. High-Level Architecture', level=1)
p('The system consists of four main layers:')
bullet('React Frontend — Login, Incident Console, Admin Portal')
bullet('FastAPI Backend — REST API with auth, incident lifecycle, admin CRUD')
bullet('LangGraph Pipeline — 12 AI agents wired into a stateful workflow')
bullet('Persistent Storage — PostgreSQL (structured data) + Qdrant (vector search)')
p('')
p('Data flows from the operator through the React frontend, to FastAPI, which '
  'invokes the LangGraph pipeline. Agents query PostgreSQL for equipment/safety/escalation '
  'rules and Qdrant for similar past incidents.')

# ── 4. Database Schema ──────────────────────────────────────────────
doc.add_heading('4. Database Schema (9 Tables)', level=1)

doc.add_heading('4.1 Core Tables', level=2)
add_table(
    ['Table', 'Purpose', 'Key Fields'],
    [
        ['users', 'System users with role-based access', 'username, email, user_type, hashed_password'],
        ['incidents', 'Core incident records + state snapshot', 'id, report, state_snapshot (JSONB), status, mttd_seconds'],
        ['qa_history', 'Question-answer exchanges per incident', 'incident_id (FK), question, answer, diagnostic_step'],
        ['outcome_records', 'Outcome attempts (success/failure)', 'incident_id (FK), outcome, resolution_summary'],
    ]
)

doc.add_heading('4.2 Operational Data Tables', level=2)
add_table(
    ['Table', 'Purpose', 'Key Fields'],
    [
        ['equipment', 'Asset registry with upstream/downstream links', 'id, name, equipment_type, criticality, upstream_id, downstream_id'],
        ['safety_rules', 'Per-type and general safety rules', 'equipment_type, rule_text, severity_class (block/constraint), is_general'],
        ['escalation_levels', 'Escalation hierarchy definitions', 'level (1-4+), name, description'],
        ['escalation_rules', 'Condition-based escalation routing', 'condition, confidence_min/max, safety_impact, escalation_level'],
        ['incident_reports', 'Historical reports for Decision Memory', 'title, asset_id, root_cause, resolution, diagnosis times'],
    ]
)

# ── 5. Application Routes ───────────────────────────────────────────
doc.add_heading('5. Application Routes', level=1)
add_table(
    ['Route', 'Component', 'Access'],
    [
        ['/login', 'Login Page', 'Public'],
        ['/*', 'Incident Console', 'Authenticated'],
        ['/admin/*', 'Admin Portal', 'Admin role only'],
    ]
)

# ── 6. Complete Incident Lifecycle ───────────────────────────────────
doc.add_heading('6. Complete Incident Lifecycle', level=1)
p('This is the core user journey — from incident report to closure.')

doc.add_heading('Phase 1: Incident Creation', level=2)
p('User describes incident → POST /api/incidents')
p('Pipeline triggered: Intake → Screening → Retrieval → Question Generation')
p('')
add_table(
    ['Agent', 'Purpose', 'Output'],
    [
        ['Intake Agent', 'Parse report → structured Incident Card', 'asset_id, severity, symptoms, safety_level, normalized_summary'],
        ['Screening Agent', 'Quick safety/risk assessment', 'risk_score (1-10), immediate escalation if safety = danger'],
        ['Retrieval Agent', 'Search Qdrant for similar past incidents', 'retrieved_patterns, retrieval_confidence, recurrence detection'],
        ['Question Agent', 'Generate diagnostic questions (10-step framework)', '1-3 questions per step, deduplicated against history'],
    ]
)

doc.add_heading('Phase 2: Diagnosis Loop (Steps 1-10)', level=2)
p('User answers question → POST /api/incidents/{id}/answer')
p('Each answer triggers: Answer Interpreter → Hypothesis Update → Safety Constraint → Advance Step → (next questions or brief)')
p('')
add_table(
    ['Agent', 'Purpose'],
    [
        ['Answer Interpreter', 'Extract facts, contradictions, and confidence signals from answer'],
        ['Hypothesis Update', 'Bayesian-style hypothesis ranking (probability 0.0 to 1.0)'],
        ['Safety Constraint', 'Check DB safety rules → blocks, constraints, warnings + audit trail'],
        ['Advance Step', 'Move to next diagnostic step (1 → 10)'],
    ]
)

doc.add_heading('The 10 Diagnostic Steps', level=3)
add_table(
    ['Step', 'Category', 'Focus'],
    [
        ['1', 'Trigger', 'What triggered the event?'],
        ['2', 'Internal Equipment', 'Equipment condition checks'],
        ['3', 'Upstream', 'Upstream process factors'],
        ['4', 'Downstream', 'Downstream restrictions'],
        ['5', 'Control System', 'Control/automation status'],
        ['6', 'Instrumentation', 'Sensor/transmitter accuracy'],
        ['7', 'Utilities', 'Supply availability'],
        ['8', 'Process Conditions', 'Operating parameter deviations'],
        ['9', 'Procedure/Human', 'Human factors'],
        ['10', 'Verification', 'Final confirmation checks'],
    ]
)

doc.add_heading('Loop Exit Conditions', level=3)
bullet('Confidence ≥ 80% (CONFIDENCE_THRESHOLD)')
bullet('Questions asked ≥ 30 (MAX_TOTAL_QUESTIONS)')
bullet('All 10 steps completed')
bullet('Escalation triggered')

doc.add_heading('Phase 3: Decision Brief', level=2)
p('Generated automatically when diagnosis exits: POST /api/incidents/{id}/brief')
add_table(
    ['Output Field', 'Description'],
    [
        ['options[]', '2-4 decision options with risk_level (low/medium/medium-high/high)'],
        ['overall_confidence', 'System confidence in the diagnosis'],
        ['decision_authority', 'Who should approve this decision'],
        ['escalation_path', 'Next person to escalate to if it fails'],
        ['risk_summary', 'Plain-language risk assessment'],
        ['safety_constraints', 'Active safety rules affecting the decision'],
    ]
)

doc.add_heading('Phase 4: Outcome Capture', level=2)
p('User reports whether the chosen action succeeded: POST /api/incidents/{id}/outcome')
bullet('Success → Verification Gate')
bullet('Failure (< 2 attempts) → Retry Diagnosis')
bullet('Failure (≥ 2 attempts) → Escalation + Expert Capture')

doc.add_heading('Phase 5: Verification & Closure', level=2)
p('POST /api/incidents/{id}/verify')
bullet('Operator confirms trigger condition has normalized')
bullet('Records verification_notes and verification_timestamp')
bullet('If verified → Memory Write Agent stores pattern to Qdrant → CLOSED')
bullet('If not verified → Escalation triggered')

doc.add_heading('Phase 6: Memory Write (Pattern Storage)', level=2)
p('Stores the complete decision pattern to Qdrant vector DB:')
add_table(
    ['Stored Field', 'Source'],
    [
        ['signals[]', 'Facts extracted during diagnosis'],
        ['root_cause', 'Confirmed root cause'],
        ['turning_point_signal', 'What signal changed the diagnosis direction'],
        ['why_symptoms_misleading', 'Expert knowledge (§11)'],
        ['escalation_rule', 'When to escalate for similar future incidents'],
        ['delay_risk', 'What happens if action is delayed'],
        ['resolution_timestamp', 'When it was resolved'],
        ['asset_id, severity', 'Incident metadata'],
    ]
)

# ── 7. Escalation Flow ──────────────────────────────────────────────
doc.add_heading('7. Escalation Flow', level=1)
p('Escalation levels and rules are fully DB-driven (admin-configurable). '
  'The escalation agent evaluates conditions in priority order:')
p('')
add_table(
    ['Priority', 'Condition', 'Level'],
    [
        ['1', 'Safety interlock / danger', 'Max level (Shutdown)'],
        ['2', 'DB escalation rules match', 'DB-determined level'],
        ['3', 'Repeated failure pattern detected', 'Level 4 (OEM)'],
        ['4', 'Risk score ≥ 8.0', 'Level 4 (OEM)'],
        ['5', '≥ 3 contradictions in diagnosis', 'Level 3 (Internal Expert)'],
        ['6', 'Confidence < 30% + previous failure', 'Level 3'],
        ['7', '≥ 2 failed attempts or high severity', 'Level 2 (Shift Engineer)'],
        ['8', 'Default (resolved at technician level)', 'Level 1 (Technician)'],
    ]
)

# ── 8. Admin Portal ─────────────────────────────────────────────────
doc.add_heading('8. Admin Portal', level=1)

doc.add_heading('8.1 Dashboard', level=2)
bullet('Total/open/closed incidents count')
bullet('Users, equipment, safety rules counts')
bullet('KPI Stats (§27): Average/Min/Max MTTD, escalation rate, process failure rate')

doc.add_heading('8.2 Management Sections', level=2)
add_table(
    ['Section', 'Features'],
    [
        ['Users', 'Create, edit, deactivate, change password, role assignment (admin/operator/engineer/viewer)'],
        ['Equipment', 'Full CRUD with upstream/downstream linking, criticality levels, process line assignment'],
        ['Safety Rules', 'Block/constraint/warning rules per equipment type, with general rules support'],
        ['Escalation Matrix', 'Level definitions + condition-based routing rules with confidence ranges'],
        ['Incidents', 'View all incidents with status, severity, confidence, and risk score'],
        ['Reports', 'Historical incident reports with Traditional vs Structured diagnosis time comparison'],
    ]
)

# Add screenshots
screenshots = [
    ('screenshot_equipment.png', 'Equipment Registry'),
    ('screenshot_safety.png', 'Safety Rules Management'),
    ('screenshot_incidents.png', 'Incidents Overview'),
    ('screenshot_escalation.png', 'Escalation Matrix Configuration'),
    ('screenshot_reports.png', 'Historical Incident Reports'),
]
artifacts_dir = '/home/code/.gemini/antigravity/brain/57cb6d12-70bb-42d4-a754-7b77e89b2d58'
for fname, caption in screenshots:
    fpath = os.path.join(artifacts_dir, fname)
    if os.path.exists(fpath):
        doc.add_heading(caption, level=3)
        doc.add_picture(fpath, width=Inches(6.5))
        doc.add_paragraph()

# ── 9. API Endpoint Summary ─────────────────────────────────────────
doc.add_heading('9. API Endpoint Summary', level=1)

doc.add_heading('9.1 Incident Lifecycle', level=2)
add_table(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/incidents', 'Create incident, run intake pipeline'],
        ['GET', '/api/incidents/{id}', 'Get incident state'],
        ['POST', '/api/incidents/{id}/answer', 'Submit diagnostic answer'],
        ['POST', '/api/incidents/{id}/brief', 'Force-generate decision brief'],
        ['POST', '/api/incidents/{id}/outcome', 'Submit outcome (success/failure)'],
        ['POST', '/api/incidents/{id}/verify', 'Verify resolution + close'],
        ['GET', '/api/incidents', 'List all incidents'],
    ]
)

doc.add_heading('9.2 Authentication', level=2)
add_table(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/auth/login', 'JWT authentication'],
        ['POST', '/api/auth/register', 'Bootstrap first admin'],
        ['GET', '/api/auth/me', 'Current user info'],
        ['PUT', '/api/auth/change-password', 'Self-service password change'],
    ]
)

doc.add_heading('9.3 Admin CRUD', level=2)
add_table(
    ['Resource', 'Endpoints'],
    [
        ['Users', 'GET / POST / PUT / DELETE  /api/admin/users'],
        ['Equipment', 'GET / POST / PUT / DELETE  /api/admin/equipment'],
        ['Safety Rules', 'GET / POST / PUT / DELETE  /api/admin/safety-rules'],
        ['Escalation Levels', 'GET / POST / PUT / DELETE  /api/admin/escalation/levels'],
        ['Escalation Rules', 'GET / POST / PUT / DELETE  /api/admin/escalation/rules'],
        ['Dashboard', 'GET  /api/admin/dashboard'],
        ['KPIs', 'GET  /api/admin/stats/kpis'],
    ]
)

doc.add_heading('9.4 Operational Data (Read-Only)', level=2)
add_table(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['GET', '/api/equipment', 'Equipment registry'],
        ['GET', '/api/safety-rules', 'Safety rules (filterable by type)'],
        ['GET', '/api/escalation-matrix', 'Escalation levels + rules'],
        ['GET', '/api/incident-reports', 'Historical reports'],
    ]
)

# ── 10. Data Flow ────────────────────────────────────────────────────
doc.add_heading('10. Data Flow Summary', level=1)
p('Complete round-trip data flow:', bold=True)
p('')
p('1. Operator reports incident via React frontend')
p('2. Frontend calls POST /api/incidents')
p('3. FastAPI runs the LangGraph intake graph (Intake → Screening → Retrieval → Questions)')
p('4. Agents query PostgreSQL for equipment info and safety rules')
p('5. Retrieval Agent searches Qdrant for similar past patterns')
p('6. State is saved to PostgreSQL, response returned to frontend')
p('7. Diagnosis loop: each answer triggers Answer → Hypothesis → Safety → Advance → Questions')
p('8. When confident enough, Decision Brief is generated with risk-rated options')
p('9. Operator executes chosen option and reports outcome')
p('10. On success: verification gate → memory write to Qdrant → incident CLOSED')
p('11. On failure: escalation agent determines level → expert capture if needed')

# ── 11. Key Design Principles ───────────────────────────────────────
doc.add_heading('11. Key Design Principles', level=1)
principles = [
    ('Database as Source of Truth', 'Equipment, safety rules, and escalation config are admin-managed in PostgreSQL'),
    ('Decision Logic, Not Repair Steps', 'System records why decisions were made, not how to fix things'),
    ('10-Step Diagnostic Framework', 'Systematic elimination of root causes across all failure categories'),
    ('Verification Gate', 'Incidents cannot close without confirming the trigger condition has normalized'),
    ('Decision Memory', 'Resolved incidents become searchable patterns for future similar incidents'),
    ('Fuzzy Deduplication', 'Questions are deduplicated against Q&A history to prevent repetition'),
    ('Audit Trail', 'Every safety rule activation is tracked with rules_applied'),
]
for title, desc in principles:
    para = doc.add_paragraph()
    run_title = para.add_run(f'{title}: ')
    run_title.bold = True
    run_title.font.size = Pt(11)
    run_desc = para.add_run(desc)
    run_desc.font.size = Pt(11)

# ── 12. File Structure ──────────────────────────────────────────────
doc.add_heading('12. Project File Structure', level=1)

add_table(
    ['Path', 'Purpose'],
    [
        ['api.py', 'FastAPI backend (1200+ lines, 30+ endpoints)'],
        ['frontend/src/App.jsx', 'React router (3 routes)'],
        ['frontend/src/LoginPage.jsx', 'JWT login page'],
        ['frontend/src/IncidentConsole.jsx', 'Main operator UI (chat-based diagnosis)'],
        ['frontend/src/AdminPortal.jsx', 'Admin dashboard + CRUD (7 sections)'],
        ['frontend/src/api.js', 'Frontend API client functions'],
        ['src/graph.py', 'LangGraph workflow (3 sub-graphs)'],
        ['src/state/state.py', 'DecisioState TypedDict + Pydantic models'],
        ['src/agents/ (12 files)', '12 LangGraph agents'],
        ['src/data/assets.py', 'Equipment lookups (DB-only)'],
        ['src/data/safety_rules.py', 'Safety rule lookups (DB-only)'],
        ['src/data/escalation_matrix.py', 'Escalation lookups (DB-only)'],
        ['src/db/models.py', 'SQLAlchemy ORM (9 tables)'],
        ['src/db/session.py', 'Async + sync session factories'],
        ['src/db/sync_queries.py', 'Sync DB queries for agents'],
        ['src/db/crud.py', 'Async CRUD operations'],
        ['src/db/seed.py', 'Database seeder'],
        ['src/llm.py', 'LLM client setup (Azure OpenAI)'],
        ['src/auth.py', 'JWT + bcrypt authentication'],
    ]
)

# ── Save ────────────────────────────────────────────────────────────
output_path = '/home/code/Decisio/Decisio_Application_Flow_Report.docx'
doc.save(output_path)
print(f'✅ Report saved to: {output_path}')
